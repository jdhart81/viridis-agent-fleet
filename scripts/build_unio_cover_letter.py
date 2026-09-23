from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "sales" / "Justin_Hart_Unio_AI_Automation_Cover_Letter_2026-08-17.docx"

NAVY = RGBColor(31, 77, 120)
INK = RGBColor(32, 36, 42)
MUTED = RGBColor(82, 89, 97)
LIGHT = "D7E0EA"


def set_font(run, size=11, bold=False, italic=False, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), LIGHT)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def paragraph(doc, text="", *, before=0, after=8, line=1.22, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    if text:
        set_font(p.add_run(text))
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.92)
    section.right_margin = Inches(0.92)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.22

    title = paragraph(doc, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(title.add_run("JUSTIN HART"), size=22, bold=True, color=NAVY)

    subtitle = paragraph(doc, after=3, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(
        subtitle.add_run("AI Systems Strategist | Agentic AI and Automation Reliability"),
        size=11.5,
        bold=True,
        color=INK,
    )

    contact = paragraph(doc, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(
        contact.add_run(
            "hartjustin6@gmail.com | 802-375-3476 | "
            "linkedin.com/in/justin-hart-639b2b269 | github.com/jdhart81"
        ),
        size=9.5,
        color=MUTED,
    )
    add_rule(contact)

    date = paragraph(doc, "August 17, 2026", after=9)
    date.runs[0].font.color.rgb = MUTED

    recipient = paragraph(doc, after=8, line=1.05)
    for i, line in enumerate(("Hiring Team", "Unio Digital", "Tucson, Arizona")):
        if i:
            recipient.add_run("\n")
        run = recipient.add_run(line)
        set_font(run, bold=(i == 0))

    subject = paragraph(doc, after=11)
    set_font(subject.add_run("Re: AI & Automation Implementation Consultant"), bold=True, color=NAVY)

    paragraph(doc, "Dear Unio Digital Hiring Team,", after=9)

    paragraph(
        doc,
        "I am applying for the AI & Automation Implementation Consultant role because it matches the work I do best: turn an ambiguous business process into a bounded implementation plan, translate executive goals into technical acceptance criteria, and carry the system through testing, deployment, monitoring, and handoff.",
    )

    paragraph(
        doc,
        "As founder of the Viridis portfolio, I design and operate API, MCP, and agent workflows with explicit approval gates, security preflight, deterministic failure handling, observability, and verifiable delivery evidence. I have led products from first-principles discovery through architecture and public deployment. A current production fleet release passes 2,169 automated tests across 37 suites, reflecting the reliability discipline I would bring to client systems rather than a demo-only approach.",
    )

    paragraph(
        doc,
        "My deepest hands-on experience is in Python, APIs, webhooks, MCP, agent orchestration, and secure deployment. I would not overstate platform-specific credentials: my n8n familiarity is workflow- and reliability-oriented rather than a claimed certification. The underlying work your role requires - discovery, integrations, retries and idempotency, approval boundaries, monitoring, client communication, and practical training - is central to how I build.",
    )

    paragraph(
        doc,
        "Unio Digital's combination of client-facing discovery, 90-day roadmaps, implementation, and team enablement is especially compelling. I would welcome the chance to discuss how my systems-thinking background and production AI experience could support your Arizona clients and project-services team in a remote contractor capacity.",
        after=12,
    )

    paragraph(doc, "Sincerely,", after=12)
    close = paragraph(doc, after=0, line=1.0)
    set_font(close.add_run("Justin Hart"), bold=True, color=NAVY)

    core = doc.core_properties
    core.title = "Justin Hart - Unio Digital AI and Automation Cover Letter"
    core.subject = "Application for AI & Automation Implementation Consultant"
    core.author = "Justin Hart"
    core.keywords = "AI automation, implementation consultant, workflow reliability, MCP, APIs"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
